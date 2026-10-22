from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUTPUT = Path("/home/nattanns18/indoor-delivery-robot/output/report/รายงาน_หัวข้อ_6_วัตถุประสงค์.docx")
FONT_NAME = "TH Sarabun New"


OBJECTIVES = [
    "เพื่อออกแบบและพัฒนาต้นแบบหุ่นยนต์เคลื่อนที่อัตโนมัติภายในอาคารที่สามารถสร้างแผนที่ ระบุตำแหน่ง วางแผนเส้นทาง และเคลื่อนที่ไปยังจุดที่กำหนดได้",
    "เพื่อพัฒนาเว็บแอปพลิเคชันสำหรับสร้างและติดตามภารกิจการเคลื่อนที่ โดยผู้ใช้งานสามารถเลือกจุดรับและจุดหมาย ตรวจสอบคิวและสถานะ ดูประวัติงานของตนเอง ยืนยันขั้นตอนของภารกิจ และยกเลิกงานที่ยังไม่เริ่มดำเนินการได้",
    "เพื่อพัฒนาระบบควบคุมส่วนกลางสำหรับให้ผู้ดูแลระบบติดตามสถานะของหุ่นยนต์ ภารกิจ แผนที่ จุดรับส่ง ภาพจากกล้อง และสถานะการทำงานของระบบผ่านเว็บแอปพลิเคชัน",
    "เพื่อพัฒนาการเชื่อมต่อระหว่างระบบเว็บกับระบบ ROS 2 ผ่าน Robot Agent สำหรับรับคำสั่งภารกิจและส่งข้อมูลตำแหน่ง เส้นทาง สถานะ และข้อมูลจากหุ่นยนต์กลับมายังระบบควบคุมส่วนกลาง",
    "เพื่อทดสอบการทำงานร่วมกันของระบบต้นแบบ ตั้งแต่การสร้างแผนที่ การกำหนดจุดรับส่ง การสร้างภารกิจ การเคลื่อนที่ระหว่างจุด และการบันทึกประวัติเหตุการณ์ภายในพื้นที่ทดลองที่กำหนด",
]


def set_run_font(run, size=16, bold=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = None
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run._element.rPr.rFonts.set(qn("w:cs"), FONT_NAME)


def set_thai_distributed(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    jc = ppr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        ppr.append(jc)
    jc.set(qn("w:val"), "thaiDistribute")


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_run_font(run, 14)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def set_top_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "4")
    top.set(qn("w:color"), "000000")
    tc_borders.append(top)
    tc_pr.append(tc_borders)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.20)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.10)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(16)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    heading = doc.add_paragraph()
    heading.paragraph_format.space_after = Pt(10)
    heading.paragraph_format.keep_with_next = True
    heading_run = heading.add_run("6.  วัตถุประสงค์")
    set_run_font(heading_run, 16, bold=True)

    for number, objective in enumerate(OBJECTIVES, start=1):
        paragraph = doc.add_paragraph()
        set_thai_distributed(paragraph)
        paragraph.paragraph_format.left_indent = Cm(1.25)
        paragraph.paragraph_format.first_line_indent = Cm(-0.63)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.widow_control = True
        number_run = paragraph.add_run(f"{number}. ")
        set_run_font(number_run, 16)
        body_run = paragraph.add_run(objective)
        set_run_font(body_run, 16)

    footer = section.footer
    footer.is_linked_to_previous = False
    table = footer.add_table(rows=1, cols=2, width=Cm(15.92))
    table.autofit = False
    table.columns[0].width = Cm(13.9)
    table.columns[1].width = Cm(2.02)
    left = table.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left_run = left.add_run("การพัฒนาหุ่นยนต์เคลื่อนที่อัตโนมัติควบคุมผ่านเว็บ")
    set_run_font(left_run, 14)
    add_page_field(table.cell(0, 1).paragraphs[0])
    for cell in table.rows[0].cells:
        set_top_border(cell)

    props = doc.core_properties
    props.title = "หัวข้อ 6 วัตถุประสงค์"
    props.subject = "การพัฒนาหุ่นยนต์เคลื่อนที่อัตโนมัติควบคุมผ่านเว็บ"
    props.author = ""
    props.keywords = "หุ่นยนต์เคลื่อนที่, เว็บแอปพลิเคชัน, ROS 2, Nav2, Robot Agent"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
