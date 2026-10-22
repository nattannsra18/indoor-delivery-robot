from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUTPUT = Path("/home/nattanns18/indoor-delivery-robot/output/report/รายงาน_หัวข้อ_5_ความสำคัญและที่มาของโครงงาน.docx")
FONT_NAME = "TH Sarabun New"


PARAGRAPHS = [
    "ในปัจจุบัน อาคารสำนักงาน ห้องปฏิบัติการ โรงงาน และคลังสินค้ามีการเคลื่อนย้ายเอกสาร อุปกรณ์ และสิ่งของขนาดเล็กระหว่างจุดต่าง ๆ ภายในอาคารอย่างต่อเนื่อง งานดังกล่าวส่วนใหญ่ยังอาศัยบุคลากรเป็นผู้ดำเนินการ ทำให้เกิดการใช้เวลาและกำลังคนกับงานที่มีลักษณะซ้ำเดิม รวมทั้งอาจติดตามสถานะของงานได้ไม่สะดวก การนำหุ่นยนต์เคลื่อนที่อัตโนมัติมาช่วยปฏิบัติงานจึงเป็นแนวทางหนึ่งที่สามารถลดภาระของบุคลากรและทำให้กระบวนการรับส่งมีสถานะที่ตรวจสอบได้ชัดเจนยิ่งขึ้น",
    "อย่างไรก็ตาม การพัฒนาหุ่นยนต์เคลื่อนที่ภายในอาคารไม่ได้ประกอบด้วยฮาร์ดแวร์และระบบนำทางเพียงอย่างเดียว แต่ยังต้องมีระบบสำหรับสร้างภารกิจ กำหนดจุดรับและจุดหมาย ติดตามสถานะ และประสานข้อมูลระหว่างผู้ใช้งาน ระบบควบคุมส่วนกลาง และหุ่นยนต์ หากการสั่งงานต้องอาศัยโปรแกรมเฉพาะทาง ผู้ใช้งานทั่วไปอาจใช้งานได้ยาก ดังนั้น การใช้เว็บแอปพลิเคชันเป็นส่วนติดต่อกับระบบจึงช่วยให้ผู้ใช้งานเข้าถึงฟังก์ชันหลักผ่านเว็บเบราว์เซอร์ได้โดยไม่จำเป็นต้องติดตั้งเครื่องมือของระบบหุ่นยนต์โดยตรง",
    "ด้วยเหตุนี้ จึงจัดทำโครงงานเรื่อง “การพัฒนาหุ่นยนต์เคลื่อนที่อัตโนมัติควบคุมผ่านเว็บ” โดยมุ่งพัฒนาต้นแบบที่เชื่อมโยงเว็บแอปพลิเคชัน ระบบควบคุมส่วนกลาง และหุ่นยนต์จริงเข้าด้วยกัน ผู้ใช้งานทั่วไปสามารถสร้างงานรับส่ง เลือกจุดรับและจุดหมาย ตรวจสอบคิวและสถานะของงาน ดูประวัติงานของตนเอง ยืนยันขั้นตอนการรับหรือส่งสิ่งของ และยกเลิกงานที่ยังไม่เริ่มดำเนินการได้ ส่วนผู้ดูแลระบบสามารถติดตามสถานะของหุ่นยนต์ ภารกิจ แผนที่ จุดรับส่ง และสถานะของระบบในภาพรวม",
    "ส่วนของหุ่นยนต์ใช้ ROS 2 และ Nav2 สำหรับการระบุตำแหน่ง การวางแผนเส้นทาง และการเคลื่อนที่บนแผนที่ภายในอาคาร โดยระบบสามารถสร้างและจัดการแผนที่ กำหนดจุดรับส่ง และส่งภารกิจจากเว็บไปยังหุ่นยนต์ผ่าน Robot Agent ซึ่งทำหน้าที่เป็นตัวกลางระหว่างระบบเว็บกับระบบ ROS 2 ขณะเดียวกัน ข้อมูลตำแหน่ง เส้นทาง สถานะภารกิจ และภาพจากกล้องจะถูกส่งกลับมาแสดงบนเว็บเพื่อให้ผู้ควบคุมติดตามการทำงานได้",
    "หุ่นยนต์ที่ใช้ในโครงงานนี้อยู่ในระดับต้นแบบฐานเคลื่อนที่ โดยติดตั้งอุปกรณ์ควบคุม อุปกรณ์ตรวจจับ และองค์ประกอบที่จำเป็นต่อการนำทาง แต่ยังไม่มีโครงสร้างสำหรับวางหรือบรรทุกสิ่งของและยังไม่มีหน้าจอสำหรับให้ผู้ใช้งานยืนยันที่ตัวหุ่นยนต์ ดังนั้น การดำเนินงานในระยะปัจจุบันจึงเน้นการพิสูจน์การทำงานของการสร้างแผนที่ การกำหนดจุดรับส่ง การระบุตำแหน่ง การสร้างภารกิจ การเคลื่อนที่ตามเส้นทาง การเปลี่ยนสถานะของงาน และการบันทึกประวัติเหตุการณ์ ภายใต้การควบคุมและการสังเกตของผู้ทดสอบ",
    "ผลที่ได้จากโครงงานมีความสำคัญในฐานะต้นแบบการบูรณาการซอฟต์แวร์เว็บ ระบบควบคุมส่วนกลาง และระบบหุ่นยนต์จริงให้ทำงานร่วมกันเป็นกระบวนการเดียว อีกทั้งยังเป็นพื้นฐานสำหรับการพัฒนาต่อในอนาคต เช่น การเพิ่มโครงสร้างสำหรับบรรทุกสิ่งของ การติดตั้งหน้าจอบนหุ่นยนต์ การยืนยันตัวตนของผู้ส่งและผู้รับ ตลอดจนการปรับปรุงความปลอดภัยและความน่าเชื่อถือเพื่อให้เหมาะสมกับการใช้งานจริงภายในอาคาร",
]


def set_run_font(run, size=16, bold=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run._element.rPr.rFonts.set(qn("w:cs"), FONT_NAME)


def set_thai_distributed(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    jc = ppr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        ppr.append(jc)
    jc.set(qn("w:val"), "thaiDistribute")


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_run_font(run, 14)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.20)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.10)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(16)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("5.  ความสำคัญและที่มาของโครงงาน")
    set_run_font(run, 16, bold=True)

    for index, text in enumerate(PARAGRAPHS):
        p = doc.add_paragraph()
        if index == 4:
            p.paragraph_format.page_break_before = True
        set_thai_distributed(p)
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.widow_control = True
        run = p.add_run(text)
        set_run_font(run, 16)

    footer = section.footer
    footer.is_linked_to_previous = False
    table = footer.add_table(rows=1, cols=2, width=Cm(15.92))
    table.autofit = False
    table.columns[0].width = Cm(13.9)
    table.columns[1].width = Cm(2.02)
    left = table.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left_run = left.add_run("การพัฒนาหุ่นยนต์เคลื่อนที่อัตโนมัติควบคุมผ่านเว็บ")
    set_run_font(left_run, 14)
    add_page_field(table.cell(0, 1).paragraphs[0])
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = OxmlElement("w:tcBorders")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "4")
        top.set(qn("w:color"), "000000")
        tc_borders.append(top)
        tc_pr.append(tc_borders)

    props = doc.core_properties
    props.title = "หัวข้อ 5 ความสำคัญและที่มาของโครงงาน"
    props.subject = "การพัฒนาหุ่นยนต์เคลื่อนที่อัตโนมัติควบคุมผ่านเว็บ"
    props.author = ""
    props.keywords = "หุ่นยนต์เคลื่อนที่, เว็บแอปพลิเคชัน, ROS 2, Nav2"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
